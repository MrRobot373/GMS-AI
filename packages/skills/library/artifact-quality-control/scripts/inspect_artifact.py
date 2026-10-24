#!/usr/bin/env python3
"""Reopen an artifact and report structural problems.

Exit 0 when nothing critical was found, 1 when critical findings exist, 2 on usage errors.
A missing inspector package is critical: the file was not verified.
`python inspect_artifact.py --self-test` checks the JSON path and, when openpyxl is
installed, the workbook error scan.
"""
from __future__ import annotations

import csv
import io
import json
import sys
import tempfile
from pathlib import Path

ERROR_MARKERS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NULL!", "#NUM!", "#GETTING_DATA")


def report_for(path: Path) -> dict:
    report: dict = {"file": path.name, "kind": path.suffix.lower().lstrip("."), "critical": [], "warnings": [], "summary": {}}
    if not path.is_file():
        report["critical"].append(f"missing file: {path}")
        return report
    ext = path.suffix.lower()
    try:
        if ext in {".xlsx", ".xlsm"}:
            inspect_xlsx(path, report)
        elif ext == ".csv":
            inspect_csv(path, report)
        elif ext == ".tsv":
            inspect_csv(path, report, delimiter="\t")
        elif ext == ".json":
            inspect_json(path, report)
        elif ext == ".pdf":
            inspect_pdf(path, report)
        elif ext == ".docx":
            inspect_docx(path, report)
        elif ext == ".pptx":
            inspect_pptx(path, report)
        else:
            report["warnings"].append(f"no inspector for {ext or 'unknown type'}; size {path.stat().st_size} bytes")
    except ModuleNotFoundError as exc:
        report["critical"].append(f"missing Python package {exc.name}; {ext} was not verified")
    return report


def inspect_xlsx(path: Path, report: dict) -> None:
    from openpyxl import load_workbook

    # Read into memory so Windows can delete the file as soon as the check ends.
    wb = load_workbook(io.BytesIO(path.read_bytes()), data_only=False)
    try:
        report["summary"]["sheets"] = wb.sheetnames
        hits: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = "" if cell.value is None else str(cell.value)
                    if any(mark in value for mark in ERROR_MARKERS):
                        hits.append(f"{ws.title}!{cell.coordinate}")
        report["summary"]["formulaErrorCells"] = len(hits)
        if hits:
            report["critical"].append("formula or cached errors in " + ", ".join(hits[:20]))
        if not wb.sheetnames:
            report["critical"].append("workbook has no sheets")
    finally:
        wb.close()


def inspect_csv(path: Path, report: dict, delimiter: str = ",") -> None:
    with path.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle, delimiter=delimiter))
    report["summary"]["rows"] = len(rows)
    if not rows:
        report["warnings"].append("delimited file has no rows")


def inspect_json(path: Path, report: dict) -> None:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        report["summary"]["items"] = len(data)
    elif isinstance(data, dict):
        report["summary"]["keys"] = len(data)
    else:
        report["summary"]["type"] = type(data).__name__


def inspect_pdf(path: Path, report: dict) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(path.read_bytes()))
    report["summary"]["pages"] = len(reader.pages)
    if not reader.pages:
        report["critical"].append("PDF has no pages")


def inspect_docx(path: Path, report: dict) -> None:
    import docx

    document = docx.Document(io.BytesIO(path.read_bytes()))
    report["summary"]["paragraphs"] = len(document.paragraphs)
    report["summary"]["tables"] = len(document.tables)


def inspect_pptx(path: Path, report: dict) -> None:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(path.read_bytes()))
    report["summary"]["slides"] = len(deck.slides)
    if not deck.slides:
        report["critical"].append("presentation has no slides")


def self_test() -> int:
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        good = folder / "sample.json"
        good.write_text('{"ok": true}\n', encoding="utf-8")
        good_report = report_for(good)
        if good_report["critical"]:
            print(json.dumps(good_report, indent=2))
            return 1
        missing = report_for(folder / "absent.json")
        if not missing["critical"]:
            print("missing file was not critical")
            return 1
        xlsx = "skipped"
        try:
            from openpyxl import Workbook
        except ModuleNotFoundError:
            xlsx = "skipped; openpyxl is not installed"
        else:
            book = Workbook()
            sheet = book.active
            sheet.title = "Checks"
            sheet["A1"] = "#DIV/0!"
            broken = folder / "broken.xlsx"
            buffer = io.BytesIO()
            book.save(buffer)
            book.close()
            broken.write_bytes(buffer.getvalue())
            broken_report = report_for(broken)
            if not broken_report["critical"]:
                print(json.dumps(broken_report, indent=2))
                return 1
            book = Workbook()
            sheet = book.active
            sheet.title = "Checks"
            sheet["A1"] = 1
            clean = folder / "clean.xlsx"
            buffer = io.BytesIO()
            book.save(buffer)
            book.close()
            clean.write_bytes(buffer.getvalue())
            clean_report = report_for(clean)
            if clean_report["critical"]:
                print(json.dumps(clean_report, indent=2))
                return 1
            xlsx = "checked"
    print(json.dumps({"selfTest": "ok", "xlsx": xlsx}))
    return 0


def main() -> int:
    if len(sys.argv) == 2 and sys.argv[1] == "--self-test":
        return self_test()
    if len(sys.argv) != 2:
        print("usage: inspect_artifact.py <file> | --self-test", file=sys.stderr)
        return 2
    report = report_for(Path(sys.argv[1]))
    print(json.dumps(report, indent=2))
    return 1 if report["critical"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
